"""
Dosya Analizi - Bağımsız Flask Uygulaması
------------------------------------------
OCR + AI destekli dosya analiz aracı.
Herhangi bir projeye bağımlı değildir.

Başlatmak için:
    python app.py
Veya:
    baslat.bat
"""

from flask import Flask, render_template, request, jsonify, send_file, send_from_directory
from werkzeug.utils import secure_filename
import os
import io
import mimetypes
import PyPDF2
import openpyxl
import pandas as pd
from docx import Document
from pptx import Presentation
from datetime import datetime

# --- Groq API helper ---
_GROQ_KEY_D = os.environ.get('GROQ_API_KEY', '')

def _call_groq(prompt, system=None, max_tokens=1500):
    import requests as _req
    if not _GROQ_KEY_D:
        return None
    try:
        msgs = [{'role':'system','content': system or 'Sen yardimci bir analist asistaninsin. Turkce cevap ver.'}]
        msgs.append({'role':'user','content':prompt})
        r = _req.post('https://api.groq.com/openai/v1/chat/completions',
            headers={'Authorization':f'Bearer {_GROQ_KEY_D}','Content-Type':'application/json'},
            json={'model':'llama-3.3-70b-versatile','messages':msgs,'max_tokens':max_tokens,'temperature':0.3},
            timeout=20)
        if r.ok:
            return r.json()['choices'][0]['message']['content'].strip()
    except Exception:
        pass
    return None


# ─────────────────────────────── KURULUM ─────────────────────────────────────

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 100 * 1024 * 1024  # 100 MB

UPLOAD_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

TESSERACT_PATH = os.environ.get('TESSERACT_PATH', '/usr/bin/tesseract')

ALLOWED_EXTENSIONS = {
    'txt', 'pdf', 'png', 'jpg', 'jpeg', 'gif', 'webp', 'bmp', 'tiff', 'tif', 'mpo',
    'doc', 'docx', 'xls', 'xlsx', 'ppt', 'pptx', 'csv',
    'mp3', 'mp4', 'avi', 'mov', 'wav',
    'json', 'xml', 'html', 'htm',
    'py', 'js', 'ts', 'css', 'md', 'yaml', 'yml', 'ini', 'toml',
    'zip', 'rar', '7z'
}

# ──────────────────────────── YARDIMCI FONKSİYONLAR ─────────────────────────

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def get_file_type(filename):
    ext = filename.rsplit('.', 1)[1].lower() if '.' in filename else ''
    if ext in ['jpg', 'jpeg', 'png', 'gif', 'bmp', 'webp', 'tiff', 'tif', 'mpo', 'heic']:
        return 'image'
    if ext == 'pdf':
        return 'pdf'
    if ext in ['doc', 'docx']:
        return 'document'
    if ext in ['xls', 'xlsx', 'csv']:
        return 'spreadsheet'
    if ext == 'pptx':
        return 'presentation'
    if ext in ['txt', 'md', 'py', 'js', 'ts', 'html', 'htm', 'css', 'json', 'xml',
               'yaml', 'yml', 'ini', 'toml']:
        return 'text'
    return 'unknown'


def format_size(size_bytes):
    for unit in ['B', 'KB', 'MB', 'GB']:
        if size_bytes < 1024:
            return f"{size_bytes:.1f} {unit}"
        size_bytes /= 1024.0
    return f"{size_bytes:.1f} TB"


def extract_text_from_path(filepath):
    """Dosya yolundan metin çıkarır. (str, str) -> (metin, hata_mesajı)"""
    filename = os.path.basename(filepath)
    ext = filename.rsplit('.', 1)[1].lower() if '.' in filename else ''
    text = ''

    try:
        # ── Metin dosyaları ──────────────────────────────────────────────────
        if ext in ['txt', 'md', 'py', 'js', 'ts', 'html', 'htm', 'css', 'json',
                   'xml', 'yaml', 'yml', 'ini', 'toml', 'csv']:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()

        # ── PDF ──────────────────────────────────────────────────────────────
        elif ext == 'pdf':
            with open(filepath, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                parts = [page.extract_text() or '' for page in reader.pages]
                text = '\n'.join(parts)

        # ── Word ─────────────────────────────────────────────────────────────
        elif ext == 'docx':
            doc = Document(filepath)
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = ' | '.join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        parts.append(row_text)
            text = '\n'.join(parts)

        # ── Excel ─────────────────────────────────────────────────────────────
        elif ext in ['xlsx', 'xls']:
            if ext == 'xlsx':
                wb = openpyxl.load_workbook(filepath, data_only=True)
                parts = []
                for sname in wb.sheetnames:
                    parts.append(f'=== {sname} ===')
                    for row in wb[sname].iter_rows(values_only=True):
                        row_str = ' | '.join(str(c) for c in row if c is not None)
                        if row_str.strip():
                            parts.append(row_str)
                text = '\n'.join(parts)
            else:
                df = pd.read_excel(filepath, sheet_name=None)
                parts = []
                for sname, sheet in df.items():
                    parts.append(f'=== {sname} ===')
                    parts.append(sheet.to_string(index=False))
                text = '\n'.join(parts)

        # ── PowerPoint ────────────────────────────────────────────────────────
        elif ext == 'pptx':
            prs = Presentation(filepath)
            parts = []
            for i, slide in enumerate(prs.slides, 1):
                parts.append(f'=== Slayt {i} ===')
                for shape in slide.shapes:
                    if hasattr(shape, 'text') and shape.text.strip():
                        parts.append(shape.text.strip())
            text = '\n'.join(parts)

        # ── Resim (OCR) ───────────────────────────────────────────────────────
        elif ext in ['jpg', 'jpeg', 'png', 'bmp', 'webp', 'tiff', 'tif', 'mpo']:
            import pytesseract
            from PIL import Image, ImageEnhance, ImageFilter

            pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH

            image = Image.open(filepath)

            # MPO (iPhone çift kamera)
            if getattr(image, 'format', '') == 'MPO':
                image.seek(0)
                tmp = filepath + '_tmp.jpg'
                image.save(tmp, 'JPEG', quality=95)
                image = Image.open(tmp)
                try:
                    os.remove(tmp)
                except:
                    pass

            if image.mode != 'RGB':
                image = image.convert('RGB')

            if max(image.size) > 3000:
                r = 3000 / max(image.size)
                image = image.resize((int(image.size[0] * r), int(image.size[1] * r)),
                                     Image.Resampling.LANCZOS)

            image = ImageEnhance.Contrast(image).enhance(1.5)
            image = ImageEnhance.Sharpness(image).enhance(2.0)
            image = image.filter(ImageFilter.MedianFilter(size=3))

            configs = [
                '--oem 3 --psm 6',
                '--oem 3 --psm 11',
                '--oem 3 --psm 4',
            ]
            best, best_score = '', 0
            for cfg in configs:
                try:
                    t = pytesseract.image_to_string(image, lang='tur+eng', config=cfg)
                    score = sum(1 for c in t if c.isalnum())
                    if score > best_score:
                        best_score, best = score, t
                except:
                    continue
            text = best

        else:
            return '', f'Desteklenmeyen dosya türü: {ext}'

    except Exception as e:
        return '', str(e)

    text = text.strip()
    if len(text) > 150_000:
        text = text[:150_000] + '\n\n... (150.000 karakterde kesildi)'

    return text, ''


# ────────────────────────────────── SAYFALAR ─────────────────────────────────

@app.route('/')
def index():
    return render_template('index.html')


@app.route('/uploads/<filename>')
def serve_upload(filename):
    """Yüklenen dosyaları serve et (resim önizlemesi için)"""
    return send_from_directory(UPLOAD_FOLDER, filename)


# ─────────────────────────────────── API ─────────────────────────────────────

@app.route('/api/upload_file', methods=['POST'])
def api_upload_file():
    try:
        if 'file' not in request.files:
            return jsonify({'success': False, 'error': 'Dosya seçilmedi'}), 400

        file = request.files['file']
        if not file or file.filename == '':
            return jsonify({'success': False, 'error': 'Dosya boş'}), 400

        if not allowed_file(file.filename):
            return jsonify({'success': False, 'error': 'Desteklenmeyen dosya türü'}), 400

        original_name = file.filename
        filename = secure_filename(original_name)
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'{timestamp}_{filename}'

        filepath = os.path.join(UPLOAD_FOLDER, filename)
        file.save(filepath)

        size = os.path.getsize(filepath)
        file_type = get_file_type(filename)

        return jsonify({
            'success': True,
            'filename': filename,
            'original_name': original_name,
            'file_path': f'/uploads/{filename}',
            'size': format_size(size),
            'file_type': file_type,
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/extract_text_from_file', methods=['POST'])
def api_extract_text_from_file():
    try:
        data = request.get_json() or {}
        filename = data.get('filename', '').strip()
        if not filename:
            return jsonify({'success': False, 'error': 'Dosya adı eksik'}), 400

        filepath = os.path.join(UPLOAD_FOLDER, secure_filename(filename))
        if not os.path.exists(filepath):
            return jsonify({'success': False, 'error': 'Dosya bulunamadı'}), 404

        text, err = extract_text_from_path(filepath)
        if err:
            return jsonify({'success': False, 'error': err})

        if not text:
            return jsonify({'success': False, 'error': 'Dosyadan metin çıkarılamadı'})

        return jsonify({'success': True, 'text': text, 'filename': filename})

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/ai_analyze_file', methods=['POST'])
def api_ai_analyze_file():
    try:
        import requests as req

        data = request.get_json() or {}
        filename     = data.get('filename', '').strip()
        analysis_type = data.get('analysis_type', 'summary')
        model        = data.get('model', 'gemma2:2b')

        if not filename:
            return jsonify({'success': False, 'error': 'Dosya adı eksik'}), 400

        filepath = os.path.join(UPLOAD_FOLDER, secure_filename(filename))
        if not os.path.exists(filepath):
            return jsonify({'success': False, 'error': 'Dosya bulunamadı'}), 404

        ext = filename.rsplit('.', 1)[1].lower() if '.' in filename else ''
        is_image = ext in ['jpg', 'jpeg', 'png', 'bmp', 'webp']
        use_groq_only = model.startswith('groq:')
        is_vision = use_groq_only or 'vision' in model.lower() or 'llava' in model.lower() or 'bakllava' in model.lower()

        # ── Prompt ──────────────────────────────────────────────────────────
        prompts = {
            'summary':     'Bu dosyanın içeriğini Türkçe olarak özetle. Önemli noktaları maddeler halinde listele.',
            'explain':     'Bu dosya ne hakkında? İçeriği detaylıca açıkla.',
            'code_review': 'Bu kodu incele. Hataları, güvenlik açıklarını ve iyileştirme önerilerini Türkçe listele.',
            'questions':   'Bu dosyayla ilgili 5 kritik soru oluştur ve cevaplarını ver.',
        }
        prompt = prompts.get(analysis_type, prompts['summary'])

        # ── Resim için base64 hazırla ────────────────────────────────────────
        img_b64 = None
        img_mime = 'image/jpeg'
        if is_image:
            import base64
            ext_lower = ext.lower()
            mime_map = {'png': 'image/png', 'gif': 'image/gif', 'webp': 'image/webp', 'bmp': 'image/bmp'}
            img_mime = mime_map.get(ext_lower, 'image/jpeg')
            with open(filepath, 'rb') as f:
                img_b64 = base64.b64encode(f.read()).decode()

        # ── Groq Vision (resim varsa — önce dene) ─────────────────────────────
        groq_vision_error = None
        if is_image and _GROQ_KEY_D:
            # Resim boyutu çok büyükse küçült
            try:
                from PIL import Image as _PILImg
                import io as _io
                with _PILImg.open(filepath) as _pil:
                    w, h = _pil.size
                    if w > 1280 or h > 1280:
                        _pil.thumbnail((1280, 1280), _PILImg.LANCZOS)
                        _buf = _io.BytesIO()
                        _pil.save(_buf, format='JPEG', quality=85)
                        img_b64 = __import__('base64').b64encode(_buf.getvalue()).decode()
                        img_mime = 'image/jpeg'
            except Exception:
                pass
            # Groq vision modelleri sırayla dene
            for _gv_model in [
                'meta-llama/llama-4-scout-17b-16e-instruct',
                'meta-llama/llama-4-maverick-17b-128e-instruct',
            ]:
                try:
                    groq_vision_resp = req.post(
                        'https://api.groq.com/openai/v1/chat/completions',
                        headers={'Authorization': f'Bearer {_GROQ_KEY_D}', 'Content-Type': 'application/json'},
                        json={
                            'model': _gv_model,
                            'messages': [{
                                'role': 'user',
                                'content': [
                                    {'type': 'text', 'text': prompt + '\n\nTürkçe cevap ver.'},
                                    {'type': 'image_url', 'image_url': {'url': f'data:{img_mime};base64,{img_b64}'}}
                                ]
                            }],
                            'max_tokens': 1500,
                            'temperature': 0.4,
                        },
                        timeout=30
                    )
                    if groq_vision_resp.ok:
                        answer = groq_vision_resp.json()['choices'][0]['message']['content'].strip()
                        return jsonify({'success': True, 'analysis': answer, 'model': f'groq:{_gv_model.split("/")[-1]}'})
                    else:
                        groq_vision_error = f'Groq ({_gv_model}): {groq_vision_resp.status_code} — {groq_vision_resp.text[:200]}'
                except Exception as _e:
                    groq_vision_error = str(_e)

        # ── Metin tabanlı Groq (resim değilse — önce dene) ───────────────────
        if not is_image and _GROQ_KEY_D:
            text_g, err_g = extract_text_from_path(filepath)
            if text_g:
                groq_ans = _call_groq(f'{prompt}\n\nDOSYA İÇERİĞİ:\n{text_g[:12000]}')
                if groq_ans:
                    return jsonify({'success': True, 'analysis': groq_ans, 'model': 'groq:llama-3.3-70b'})
            if use_groq_only:
                return jsonify({'success': False, 'error': 'Groq API yanıt vermedi. Lütfen birkaç saniye sonra tekrar deneyin.'})

        # ── Resim + vision model (Ollama) ────────────────────────────────────
        if is_image and is_vision:
            payload = {
                'model': model,
                'prompt': prompt,
                'images': [img_b64],
                'stream': False,
            }
            timeout = 120
        elif is_image:
            # Resim ama Groq Vision başarısız + Ollama vision yok → hata ver, OCR yapma
            err_detail = f' (Detay: {groq_vision_error})' if groq_vision_error else ''
            return jsonify({'success': False,
                            'error': f'Resim analizi yapılamadı. Groq Vision API\'ye erişilemedi{err_detail}. Lütfen internet bağlantınızı kontrol edin veya birkaç saniye sonra tekrar deneyin.'})
        else:
            # Metin çıkar
            text, err = extract_text_from_path(filepath)
            if err:
                return jsonify({'success': False, 'error': err})
            if not text:
                return jsonify({'success': False, 'error': 'Dosyadan metin çıkarılamadı'})

            full_prompt = f'{prompt}\n\nDOSYA İÇERİĞİ:\n{text[:15000]}'
            payload = {
                'model': model,
                'prompt': full_prompt,
                'stream': False,
                'options': {'temperature': 0.3, 'num_predict': 1024},
            }
            timeout = 120

        # ── Ollama çağrısı ────────────────────────────────────────────────────
        try:
            resp = req.post('http://localhost:11434/api/generate',
                            json=payload, timeout=timeout)
        except Exception as conn_err:
            # Ollama yok — Groq metin fallback (henüz denenmemişse)
            if not is_image and _GROQ_KEY_D:
                try:
                    text_fb, _ = extract_text_from_path(filepath)
                    if text_fb:
                        groq_ans = _call_groq(f'{prompt}\n\nDOSYA İÇERİĞİ:\n{text_fb[:12000]}')
                        if groq_ans:
                            return jsonify({'success': True, 'analysis': groq_ans, 'model': 'groq:llama-3.3-70b'})
                except Exception:
                    pass
            return jsonify({'success': False,
                            'error': f'Ollama bağlantı hatası: {conn_err}'})

        if resp.status_code == 404:
            # Model bulunamadı → küçük bir fallback dene
            fallback = 'gemma2:2b' if not is_vision else 'llava:7b'
            payload['model'] = fallback
            try:
                resp = req.post('http://localhost:11434/api/generate',
                                json=payload, timeout=timeout)
                if resp.status_code == 200:
                    model = fallback
            except:
                pass

        if resp.status_code == 200:
            result = resp.json()
            return jsonify({
                'success': True,
                'analysis': result.get('response', '').strip(),
                'model': model,
            })

        return jsonify({'success': False,
                        'error': f'Ollama hatası: {resp.status_code} — {resp.text[:300]}'})

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/extract_text', methods=['POST'])
def api_extract_text_upload():
    """Doğrudan dosya upload ile OCR (form-data)"""
    try:
        if 'file' not in request.files and 'image' not in request.files:
            return jsonify({'success': False, 'error': 'Dosya seçilmedi'}), 400

        file = request.files.get('file') or request.files.get('image')
        if not file or file.filename == '':
            return jsonify({'success': False, 'error': 'Dosya boş'}), 400

        import pytesseract
        from PIL import Image
        pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH

        img = Image.open(file.stream)
        text = pytesseract.image_to_string(img, lang='tur+eng')

        return jsonify({'success': True, 'extracted_text': text, 'text': text})

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/ollama_modeller')
def api_ollama_models():
    try:
        import requests as req
        resp = req.get('http://localhost:11434/api/tags', timeout=3)
        if resp.status_code == 200:
            models = [m['name'] for m in resp.json().get('models', [])]
            return jsonify({
                'success': True,
                'models': models,
                'recommended': models[0] if models else 'gemma2:2b',
            })
    except:
        pass

    # Ollama yoksa fallback
    fallback = ['gemma2:2b', 'llama3.2', 'llama3.2-vision', 'llava:7b',
                'qwen2.5:7b', 'mistral']
    return jsonify({
        'success': True,
        'models': fallback,
        'recommended': fallback[0],
        'warning': 'Ollama servisine erişilemedi, varsayılan liste gösteriliyor',
    })


@app.route('/api/translate', methods=['POST'])
def api_translate():
    try:
        data = request.get_json() or {}
        text        = data.get('text', '')
        source_lang = data.get('source_lang', 'tr')
        target_lang = data.get('target_lang', 'en')

        if not text.strip():
            return jsonify({'success': False, 'error': 'Metin boş'})

        from deep_translator import GoogleTranslator

        MAX = 4500
        if len(text) <= MAX:
            translated = GoogleTranslator(source=source_lang,
                                          target=target_lang).translate(text)
        else:
            parts = []
            for i in range(0, len(text), MAX):
                chunk = text[i:i + MAX]
                parts.append(GoogleTranslator(source=source_lang,
                                              target=target_lang).translate(chunk))
            translated = '\n'.join(parts)

        return jsonify({
            'success': True,
            'translated_text': translated,
            'source_lang': source_lang,
            'target_lang': target_lang,
        })

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})


@app.route('/api/delete_file/<filename>', methods=['DELETE'])
def api_delete_file(filename):
    try:
        filepath = os.path.join(UPLOAD_FOLDER, secure_filename(filename))
        if os.path.exists(filepath):
            os.remove(filepath)
            return jsonify({'success': True})
        return jsonify({'success': False, 'error': 'Dosya bulunamadı'}), 404
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


# ─────────────────────────────────── BAŞLAT ──────────────────────────────────


# ── PWA ──────────────────────────────────────────
@app.route('/manifest.json')
def pwa_manifest():
    return send_from_directory(app.static_folder, 'manifest.json',
                               mimetype='application/manifest+json')

@app.route('/sw.js')
def pwa_sw():
    resp = send_from_directory(app.static_folder, 'sw.js',
                               mimetype='application/javascript')
    resp.headers['Service-Worker-Allowed'] = '/'
    return resp
# ─────────────────────────────────────────────────

if __name__ == '__main__':
    print("\n" + "=" * 55)
    print("  Dosya Analizi Uygulaması")
    print("  http://localhost:5050")
    print("=" * 55 + "\n")
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 5050)), debug=False)
